import os
import re
import tempfile
import threading
import traceback
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from difflib import SequenceMatcher
from collections import OrderedDict

from docx import Document
from docx.oxml.ns import qn
from docx.table import _Cell
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from pyexcel_ods3 import save_data, get_data

from odf.opendocument import load as odf_load
from odf.table import Table, TableRow, TableCell
from odf.text import P

try:
    import xlwt
    _HAS_XLWT = True
except ImportError:
    _HAS_XLWT = False


# ============================================================
# НАСТРОЙКИ
# ============================================================
IDLE_TIMEOUT_SEC = 15 * 60   # 15 минут


# ============================================================
# ШАПКИ ФОРМАТОВ
# ============================================================
FORMAT_32 = [
    ("Номер формата", "AUTO"),
    ("Код СФО", None),
    ("Код НОПС", 0),
    ("НОПС в единственном числе", 1),
    ("НОПС во множественном числе", None),
    ("Класс ЕКПС", 2),
    ("Номер позиции базовых характеристик", 4),
    ("Номер характеристики обязательной к заполнению", 5),
]

FORMAT_35 = [
    ("Номер формата", "AUTO"),
    ("ФНН ПС", ["фнн"]),
    ("Номер позиции", ["поз"]),
    ("Номер измерения", ["измер"]),
    ("Код параметра", ["код", "параметр"]),
    ("Код объекта", ["код", "объект"]),
    ("Наименование характеристики", ["наимен", "характер"]),
    ("Функция соответствия", ["функц", "соответств"]),
    ("Код единицы измерения", ["единиц", "измер", "код"]),
    ("Полное наименование единицы измерения", ["единиц", "измер", "наимен"]),
    ("Сокращенное наименование единицы измерения", ["сокращ", "единиц"]),
    ("Код значения", ["значен", "код"]),
    ("Полное наименование значения", ["значен", "наимен"]),
    ("Сокращенное наименование значения", ["сокращ", "значен"]),
]

FORMAT_38 = [
    ("Номер формата", "AUTO"),
    ("Код СФО", None),
    ("Номер позиции", ["поз"]),
    ("Тип характеристики", ["вид", "характер"]),
    ("Код параметра", ["код", "параметр"]),
    ("Код объекта", ["код", "объект"]),
    ("Падеж объекта", ["падеж"]),
    ("Наименование характеристики (параметра и объекта)", ["наимен", "характер"]),
    ("Обозначение функции соответствия", ["функц", "соответств"]),
    ("Тип измеряемой величины", ["тип", "измер"]),
    ("Код единицы измерения", ["единиц", "измер", "код"]),
    ("Полное наименование единицы измерения", ["единиц", "измер", "наимен"]),
    ("Сокращенное наименование единицы измерения", ["сокращ", "единиц"]),
    ("Код возможного значения условия", ["возможн", "услов", "код"]),
    ("Полное наименование возможного значения условия", ["возможн", "услов", "наимен"]),
    ("Сокращенное наименование возможного значения условия", ["сокращ", "возможн"]),
    ("Код характера изменения свойства", ["характер", "изменен", "код"]),
    ("Нижняя граница значения характера изменения свойства", ["нижн", "границ"]),
    ("Верхняя граница значения характера изменения свойства", ["верхн", "границ"]),
    ("Номер блока характеристик", ["блок", "характер"]),
]

FORMATS = {"32": FORMAT_32, "35": FORMAT_35, "38": FORMAT_38}

SFO_COLUMN_INDEX = {"32": 1, "38": 1}
FUNC_MATCH_COLUMNS = {"35": [7], "38": [8]}
FUZZY_THRESHOLD = 0.75

EXT_MAP = {
    "xlsx": ".xlsx",
    "xlsm": ".xlsm",
    "xls":  ".xls",
    "ods":  ".ods",
}


# ============================================================
# ОПРЕДЕЛЕНИЕ ТИПА ФАЙЛА
# ============================================================
def _file_signature(path, n=8):
    with open(path, "rb") as f:
        return f.read(n)


def _is_real_doc(path):
    try:
        return _file_signature(path, 8) == b"\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1"
    except Exception:
        return False


def _is_zip(path):
    try:
        return _file_signature(path, 4) == b"PK\x03\x04"
    except Exception:
        return False


def _zip_mimetype(path):
    import zipfile
    try:
        with zipfile.ZipFile(path) as z:
            if "mimetype" in z.namelist():
                return z.read("mimetype").decode("ascii", errors="ignore").strip()
            if "word/document.xml" in z.namelist():
                return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
            if "xl/workbook.xml" in z.namelist():
                return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    except Exception:
        pass
    return ""


def detect_source_kind(path):
    ext = os.path.splitext(path)[1].lower()

    if _is_real_doc(path):
        return "doc"

    if _is_zip(path):
        mime = _zip_mimetype(path)
        if "opendocument.text" in mime:
            return "odt"
        if "opendocument.spreadsheet" in mime:
            return "ods"
        if "wordprocessingml" in mime:
            return "docx"
        if "spreadsheetml" in mime:
            return "docx"
        if ext == ".odt":
            return "odt"
        if ext == ".ods":
            return "ods"
        return "docx"

    if ext == ".odt":
        return "odt"
    if ext == ".ods":
        return "ods"
    return "unknown"


def ensure_docx(path):
    kind = detect_source_kind(path)
    if kind != "doc":
        return path

    try:
        import aspose.words_foss as aw
    except ImportError:
        raise RuntimeError(
            "Для чтения .doc нужен пакет aspose-words-foss.\n"
            "Установите: python3.12 -m pip install --user --break-system-packages aspose-words-foss"
        )

    tmp_dir = tempfile.mkdtemp(prefix="doc_convert_")
    base = os.path.splitext(os.path.basename(path))[0]
    out_docx = os.path.join(tmp_dir, base + ".docx")

    doc = aw.Document(path)
    doc.save(out_docx, aw.SaveFormat.DOCX)
    if not os.path.exists(out_docx):
        raise RuntimeError(f"Не удалось сконвертировать {path} в .docx")
    return out_docx


# ============================================================
# ЛОГИКА
# ============================================================
def normalize(text):
    if text is None:
        return ""
    text = str(text).strip().lower().replace("\n", " ")
    text = re.sub(r"[«»\"'`]", "", text)
    text = re.sub(r"[.:;,!?()\[\]{}]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def similarity(a, b):
    a, b = normalize(a), normalize(b)
    if not a or not b:
        return 0.0
    if a == b:
        return 1.0
    if a in b or b in a:
        return 0.95
    return SequenceMatcher(None, a, b).ratio()


def has_all_markers(text, markers):
    if not markers:
        return True
    t = normalize(text)
    t_nospace = re.sub(r"\s+", "", t)
    return all((m in t) or (re.sub(r"\s+", "", m) in t_nospace) for m in markers)


# ---------- ЧТЕНИЕ .docx ----------
def clean_cell_text(cell):
    paragraphs = [p.text.strip() for p in cell.paragraphs]
    paragraphs = [p for p in paragraphs if p]
    seen = set()
    result = []
    for p in paragraphs:
        key = re.sub(r"[^\w]+", "", p.lower(), flags=re.UNICODE)
        if key:
            if key not in seen:
                seen.add(key)
                result.append(p)
        else:
            result.append(p)
    return " ".join(result).strip()


def get_grid_span(tc):
    tcPr = tc.find(qn("w:tcPr"))
    if tcPr is None:
        return 1
    gs = tcPr.find(qn("w:gridSpan"))
    if gs is None:
        return 1
    return int(gs.get(qn("w:val"), "1"))


def get_vmerge(tc):
    tcPr = tc.find(qn("w:tcPr"))
    if tcPr is None:
        return None
    vm = tcPr.find(qn("w:vMerge"))
    if vm is None:
        return None
    val = vm.get(qn("w:val"))
    return val if val else "continue"


def read_word_table(table):
    n_cols = len(table.columns)
    vmerge_values = [None] * n_cols
    result = []
    for row in table.rows:
        row_values = [None] * n_cols
        col_idx = 0
        for tc in row._tr.findall(qn("w:tc")):
            if col_idx >= n_cols:
                break
            span = get_grid_span(tc)
            vmerge = get_vmerge(tc)
            if vmerge == "continue":
                value = vmerge_values[col_idx]
            else:
                cell_obj = _Cell(tc, table)
                value = clean_cell_text(cell_obj)
                if vmerge == "restart":
                    vmerge_values[col_idx] = value
                else:
                    vmerge_values[col_idx] = None
            for _ in range(span):
                if col_idx < n_cols:
                    row_values[col_idx] = value
                    col_idx += 1
        while col_idx < n_cols:
            row_values[col_idx] = vmerge_values[col_idx]
            col_idx += 1
        result.append(row_values)
    return result


def read_docx_tables(path):
    doc = Document(path)
    all_rows = []
    for t in doc.tables:
        all_rows.extend(read_word_table(t))
    return all_rows


# ---------- ЧТЕНИЕ .odt ----------
def _odf_get_text(cell):
    parts = []
    for p in cell.getElementsByType(P):
        txt = " ".join(str(n) for n in p.childNodes if n.nodeType == n.TEXT_NODE)
        if not txt:
            txt = "".join(
                getattr(child, "data", "") for child in p.childNodes
                if getattr(child, "nodeType", None) == child.TEXT_NODE
            )
        txt = txt.strip()
        if txt:
            parts.append(txt)
    seen = set()
    out = []
    for p in parts:
        if p not in seen:
            seen.add(p)
            out.append(p)
    return " ".join(out).strip()


def _odf_cell_repeat(cell):
    r = cell.getAttribute("numbercolumnsrepeated")
    try:
        return int(r) if r else 1
    except Exception:
        return 1


def _odf_row_repeat(row):
    r = row.getAttribute("numberrowsrepeated")
    try:
        return int(r) if r else 1
    except Exception:
        return 1


def read_odt_tables(path):
    odt = odf_load(path)
    all_rows = []
    for table in odt.getElementsByType(Table):
        for row in table.getElementsByType(TableRow):
            row_repeat = _odf_row_repeat(row)
            cells = row.getElementsByType(TableCell)
            row_values = []
            for c in cells:
                txt = _odf_get_text(c)
                rep = _odf_cell_repeat(c)
                for _ in range(rep):
                    row_values.append(txt)
            for _ in range(row_repeat):
                all_rows.append(list(row_values))
    return all_rows


# ---------- ЧТЕНИЕ .ods ----------
def read_ods_tables(path):
    data = get_data(path)
    all_rows = []
    for sheet_name, rows in data.items():
        for r in rows:
            row_values = []
            for v in r:
                if v is None:
                    row_values.append("")
                elif isinstance(v, float) and v.is_integer():
                    row_values.append(str(int(v)))
                else:
                    row_values.append(str(v))
            all_rows.append(row_values)
    return all_rows


# ---------- УНИВЕРСАЛЬНОЕ ЧТЕНИЕ ----------
def read_source_tables(path):
    kind = detect_source_kind(path)
    if kind == "doc":
        docx_path = ensure_docx(path)
        return read_docx_tables(docx_path)
    if kind == "docx":
        return read_docx_tables(path)
    if kind == "odt":
        return read_odt_tables(path)
    if kind == "ods":
        return read_ods_tables(path)
    return read_docx_tables(path)


# ---------- ОПРЕДЕЛЕНИЕ ШАПКИ ----------
def is_number_row(row):
    if not row:
        return False
    non_empty = [c for c in row if c and str(c).strip()]
    if not non_empty:
        return False
    nums = sum(1 for c in non_empty if re.fullmatch(r"\d+", str(c).strip()))
    return nums >= len(non_empty) * 0.7


def detect_word_header(rows):
    markers = ["наимен", "код", "номер", "значен", "тип", "единица",
               "характер", "функц", "измер", "поз", "вид", "объект",
               "нопс", "екпс", "класс", "описание"]

    best_idx, best_score = 0, -1
    for i, row in enumerate(rows[:10]):
        joined = normalize(" ".join(str(c) for c in row if c))
        joined_ns = re.sub(r"\s+", "", joined)
        score = sum(1 for m in markers if (m in joined) or (m in joined_ns))
        numeric = sum(1 for c in row if c and re.fullmatch(r"\d+", str(c).strip()))
        score -= numeric * 0.5
        if score > best_score:
            best_score, best_idx = score, i

    top = rows[best_idx]
    bottom = rows[best_idx + 1] if best_idx + 1 < len(rows) else None

    if bottom is not None and is_number_row(bottom):
        merged = list(top)
        data_start = best_idx + 2
    else:
        if bottom is not None and not any(str(c).strip() for c in bottom if c):
            merged = list(top)
            data_start = best_idx + 1
        else:
            merged = []
            n = max(len(top), len(bottom) if bottom else 0)
            for i in range(n):
                t = top[i] if i < len(top) else ""
                b = bottom[i] if bottom and i < len(bottom) else ""
                t_str = str(t).strip() if t else ""
                b_str = str(b).strip() if b else ""
                if t_str and b_str and t_str != b_str:
                    merged.append(f"{t_str} {b_str}")
                elif t_str:
                    merged.append(t_str)
                elif b_str:
                    merged.append(b_str)
                else:
                    merged.append("")
            data_start = best_idx + 2

    if data_start < len(rows) and is_number_row(rows[data_start]):
        data_start += 1

    return best_idx, merged, rows[data_start:]


def build_mapping(word_headers, fmt_columns):
    hard_mode = True
    for name, val in fmt_columns:
        if isinstance(val, list):
            hard_mode = False
            break

    mapping = []
    if hard_mode:
        for name, val in fmt_columns:
            mapping.append("AUTO" if val == "AUTO" else val)
        return mapping

    used = set()
    for fmt_name, markers in fmt_columns:
        if markers == "AUTO":
            mapping.append("AUTO")
            continue
        if markers is None:
            mapping.append(None)
            continue

        best_idx, best_score = None, 0.0
        for w_idx, w_name in enumerate(word_headers):
            if w_idx in used:
                continue
            if not w_name:
                continue
            if not has_all_markers(w_name, markers):
                continue
            s = similarity(fmt_name, w_name)
            if s > best_score:
                best_score, best_idx = s, w_idx

        if best_idx is not None and best_score >= FUZZY_THRESHOLD:
            mapping.append(best_idx)
            used.add(best_idx)
        else:
            mapping.append(None)

    return mapping


def unique_path(folder, base_name, ext=".xlsx"):
    path = os.path.join(folder, f"{base_name}{ext}")
    if not os.path.exists(path):
        return path
    i = 1
    while True:
        path = os.path.join(folder, f"{base_name} ({i}){ext}")
        if not os.path.exists(path):
            return path
        i += 1


# ============================================================
# СОХРАНЕНИЕ
# ============================================================
def _save_xlsx(out_path, fmt_columns, excel_rows, out_format):
    wb = Workbook()
    ws = wb.active
    ws.title = out_format

    bold = Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col_idx, (name, _) in enumerate(fmt_columns, start=1):
        c = ws.cell(row=1, column=col_idx, value=name)
        c.font = bold
        c.alignment = center
        c.border = border

    for r_off, row_data in enumerate(excel_rows):
        excel_row = 2 + r_off
        for col_idx, value in enumerate(row_data, start=1):
            c = ws.cell(row=excel_row, column=col_idx, value=value)
            c.alignment = left
            c.border = border

    for col_idx in range(1, len(fmt_columns) + 1):
        max_len = 0
        for row_idx in range(1, ws.max_row + 1):
            v = ws.cell(row=row_idx, column=col_idx).value
            if v is not None:
                max_len = max(max_len, len(str(v)))
        width = min(max(max_len + 2, 10), 50)
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    wb.save(out_path)


def _save_xls(out_path, fmt_columns, excel_rows):
    if not _HAS_XLWT:
        raise RuntimeError(
            "Для сохранения в .xls нужен пакет xlwt.\n"
            "Установите: python3.12 -m pip install --user --break-system-packages xlwt"
        )

    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("Sheet1")

    bold_font = xlwt.Font()
    bold_font.bold = True
    header_style = xlwt.XFStyle()
    header_style.font = bold_font
    header_style.alignment = xlwt.Alignment()
    header_style.alignment.horz = xlwt.Alignment.HORZ_CENTER
    header_style.alignment.vert = xlwt.Alignment.VERT_CENTER
    header_style.alignment.wrap = 1

    body_style = xlwt.XFStyle()
    body_style.alignment = xlwt.Alignment()
    body_style.alignment.horz = xlwt.Alignment.HORZ_LEFT
    body_style.alignment.vert = xlwt.Alignment.VERT_CENTER
    body_style.alignment.wrap = 1

    for col_idx, (name, _) in enumerate(fmt_columns):
        ws.write(0, col_idx, name, header_style)

    for r_off, row_data in enumerate(excel_rows):
        excel_row = 1 + r_off
        for col_idx, value in enumerate(row_data):
            if value is None:
                value = ""
            ws.write(excel_row, col_idx, value, body_style)

    for col_idx in range(len(fmt_columns)):
        max_len = 0
        for row_idx in range(len(excel_rows)):
            v = excel_rows[row_idx][col_idx] if col_idx < len(excel_rows[row_idx]) else None
            if v is not None:
                max_len = max(max_len, len(str(v)))
        width = min(max(max_len + 2, 10), 50)
        ws.col(col_idx).width = width * 256

    wb.save(out_path)


def _save_ods(out_path, fmt_columns, excel_rows):
    data = OrderedDict()
    sheet_data = []

    header = [name for name, _ in fmt_columns]
    sheet_data.append(header)

    for row_data in excel_rows:
        clean_row = ["" if v is None else v for v in row_data]
        sheet_data.append(clean_row)

    data["Sheet1"] = sheet_data
    save_data(out_path, data)


# ============================================================
# ОСНОВНАЯ ФУНКЦИЯ ПЕРЕНОСА
# ============================================================
def do_transfer(word_path, format_key, out_format="xlsx",
                out_path_manual=None, sfo_code=""):
    raw = read_source_tables(word_path)
    if not raw:
        raise ValueError("В исходном файле не найдено ни одной таблицы.")

    fmt_columns = FORMATS[format_key]
    idx, w_head, w_data = detect_word_header(raw)

    print("=" * 70)
    print(f"Формат: {format_key}. Шапка в строке №{idx + 1} (всего строк: {len(raw)})")
    print("=" * 70)
    print("\nШАПКА:")
    for i, h in enumerate(w_head):
        print(f"  [{i}] {repr(h)}")
    print("\nПервые 3 строки данных:")
    for r_i, r in enumerate(w_data[:3]):
        print(f"  --- строка {r_i + 1} (всего ячеек: {len(r)}) ---")
        for c_i, c in enumerate(r):
            print(f"    [{c_i}] {repr(c)}")
    print("=" * 70)

    if not w_data:
        raise ValueError("В таблице нет строк с данными.")

    mapping = build_mapping(w_head, fmt_columns)

    print("\nСОПОСТАВЛЕНИЕ КОЛОНОК:")
    for col_idx, ((fmt_name, _), w_idx) in enumerate(zip(fmt_columns, mapping)):
        if w_idx == "AUTO":
            print(f"  ★ [{col_idx}] {fmt_name}  — АВТОЗАПОЛНЕНИЕ")
        elif w_idx is None:
            print(f"  ✗ [{col_idx}] {fmt_name}  — НЕ НАЙДЕНО (пусто)")
        else:
            print(f"  ✓ [{col_idx}] {fmt_name}  ←  Word[{w_idx}] = {repr(w_head[w_idx])}")
    print("=" * 70)

    sfo_idx = SFO_COLUMN_INDEX.get(format_key)
    func_out_cols = set(FUNC_MATCH_COLUMNS.get(format_key, []))

    excel_rows = []
    for wrow in w_data:
        row_data = [None] * len(fmt_columns)
        for col_pos, w_idx in enumerate(mapping):
            if sfo_idx is not None and col_pos == sfo_idx:
                row_data[col_pos] = sfo_code if sfo_code else None
                continue
            if w_idx == "AUTO":
                row_data[col_pos] = format_key
                continue
            if w_idx is None:
                row_data[col_pos] = None
                continue
            value = wrow[w_idx] if w_idx < len(wrow) else None
            if value == "":
                value = None
            row_data[col_pos] = value

        for func_col in func_out_cols:
            if func_col >= len(row_data):
                continue
            v = row_data[func_col]
            if v:
                s = str(v).strip()
                row_data[func_col] = "=" if (s == "=" or s.startswith("=")) else None
            else:
                row_data[func_col] = None

        excel_rows.append(row_data)

    if out_path_manual:
        folder = os.path.dirname(os.path.abspath(out_path_manual)) or "."
        os.makedirs(folder, exist_ok=True)
        base = os.path.splitext(os.path.basename(out_path_manual))[0]
        ext = os.path.splitext(out_path_manual)[1] or EXT_MAP.get(out_format, ".xlsx")
        out_path = unique_path(folder, base, ext=ext)
    else:
        folder = os.path.dirname(os.path.abspath(word_path))
        ext = EXT_MAP.get(out_format, ".xlsx")
        out_path = unique_path(folder, format_key, ext=ext)

    if out_format == "ods":
        _save_ods(out_path, fmt_columns, excel_rows)
    elif out_format == "xls":
        _save_xls(out_path, fmt_columns, excel_rows)
    else:
        _save_xlsx(out_path, fmt_columns, excel_rows, out_format)

    return out_path


# ============================================================
# GUI
# ============================================================
class App:
    def __init__(self, root):
        self.root = root
        root.title("Каталогизация-форматы")
        root.geometry("760x620")
        root.minsize(720, 600)
        root.resizable(False, False)

        self.word_path = tk.StringVar()
        self.out_path_var = tk.StringVar()
        self.out_format_var = tk.StringVar(value="xlsx")
        self.sfo_var = tk.StringVar(value="")

        self._save_folder = None

        self._idle_job = None
        self._idle_paused = False

        container = tk.Frame(root)
        container.pack(fill="both", expand=True, padx=14, pady=14)

        tk.Label(container, text="Код СФО:", anchor="w",
                 font=("Arial", 11, "bold")).pack(fill="x", pady=(0, 4))
        f0 = tk.Frame(container)
        f0.pack(fill="x", pady=(0, 12))
        tk.Entry(f0, textvariable=self.sfo_var, width=30).pack(side="left")

        tk.Label(container, text="Исходный файл:", anchor="w",
                 font=("Arial", 11, "bold")).pack(fill="x", pady=(0, 4))
        f1 = tk.Frame(container)
        f1.pack(fill="x", pady=(0, 12))
        tk.Entry(f1, textvariable=self.word_path).pack(side="left", fill="x", expand=True)
        tk.Button(f1, text="Выбрать…", width=12, command=self.pick_word).pack(side="left", padx=(6, 0))

        tk.Label(container, text="Место сохранения:", anchor="w",
                 font=("Arial", 11, "bold")).pack(fill="x", pady=(0, 4))
        f2 = tk.Frame(container)
        f2.pack(fill="x", pady=(0, 12))
        tk.Entry(f2, textvariable=self.out_path_var).pack(side="left", fill="x", expand=True)
        tk.Button(f2, text="Выбрать…", width=12, command=self.pick_save_path).pack(side="left", padx=(6, 0))

        tk.Label(container, text="Формат сохранения:", anchor="w",
                 font=("Arial", 11, "bold")).pack(fill="x", pady=(0, 4))
        f3 = tk.Frame(container)
        f3.pack(fill="x", pady=(0, 6))
        self.out_format_menu = tk.OptionMenu(f3, self.out_format_var, "xlsx", "xls", "xlsm", "ods",
                                             command=lambda _=None: self._on_format_changed())
        self.out_format_menu.config(width=20, anchor="w")
        self.out_format_menu.pack(side="left")

        tk.Label(
            container,
            text=(
                "xlsx — современный Excel.\n"
                "xls — старый Excel.\n"
                "xlsm — Excel с макросами.\n"
                "ods — LibreOffice / OpenOffice."
            ),
            anchor="w", justify="left",
            fg="#666", font=("Arial", 9),
        ).pack(fill="x", pady=(0, 14))

        tk.Label(container, text="Выберите форму:", anchor="w",
                 font=("Arial", 11, "bold")).pack(fill="x", pady=(0, 6))
        f4 = tk.Frame(container)
        f4.pack(fill="x")

        btn_style = {
            "font": ("Arial", 12, "bold"),
            "height": 2,
            "cursor": "hand2",
            "bg": "#4CAF50",
            "fg": "white",
            "activebackground": "#45a049",
            "activeforeground": "white",
            "relief": "raised",
            "bd": 2,
        }
        self.btn32 = tk.Button(f4, text="Форма 32", command=lambda: self.run("32"), **btn_style)
        self.btn35 = tk.Button(f4, text="Форма 35", command=lambda: self.run("35"), **btn_style)
        self.btn38 = tk.Button(f4, text="Форма 38", command=lambda: self.run("38"), **btn_style)
        self.btn32.pack(side="left", fill="x", expand=True, padx=(0, 4))
        self.btn35.pack(side="left", fill="x", expand=True, padx=4)
        self.btn38.pack(side="left", fill="x", expand=True, padx=(4, 0))
        self.form_buttons = [self.btn32, self.btn35, self.btn38]

        self.word_path.trace_add("write", self._on_word_changed)

        self._bind_idle_events()
        self._reset_idle_timer()

    def _bind_idle_events(self):
        r = self.root
        r.bind_all("<Any-KeyPress>", self._on_user_activity, add="+")
        r.bind_all("<Any-Button>", self._on_user_activity, add="+")
        r.bind_all("<Motion>", self._on_user_activity, add="+")
        r.bind_all("<MouseWheel>", self._on_user_activity, add="+")

    def _on_user_activity(self, _event=None):
        if self._idle_paused:
            return
        self._reset_idle_timer()

    def _reset_idle_timer(self):
        if self._idle_job is not None:
            try:
                self.root.after_cancel(self._idle_job)
            except Exception:
                pass
            self._idle_job = None
        if self._idle_paused:
            return
        self._idle_job = self.root.after(IDLE_TIMEOUT_SEC * 1000, self._on_idle_timeout)

    def _on_idle_timeout(self):
        try:
            self.root.destroy()
        except Exception:
            pass

    def _pause_idle(self):
        self._idle_paused = True
        if self._idle_job is not None:
            try:
                self.root.after_cancel(self._idle_job)
            except Exception:
                pass
            self._idle_job = None

    def _resume_idle(self):
        self._idle_paused = False
        self._reset_idle_timer()

    def _current_folder(self):
        if self._save_folder:
            return self._save_folder
        w = self.word_path.get().strip()
        if w:
            return os.path.dirname(os.path.abspath(w))
        return ""

    def _show_folder_only(self):
        folder = self._current_folder()
        self.out_path_var.set(folder if folder else "")

    def _show_full_target(self, fmt):
        folder = self._current_folder()
        if not folder:
            return ""
        ext = EXT_MAP.get(self.out_format_var.get(), ".xlsx")
        full = unique_path(folder, fmt, ext=ext)
        self.out_path_var.set(full)
        return full

    def _on_word_changed(self, *_):
        w = self.word_path.get().strip()
        if not w:
            return
        if not self._save_folder:
            self._show_folder_only()

    def _on_format_changed(self):
        raw = self.out_path_var.get().strip()
        if not raw:
            return
        if os.path.isdir(raw):
            self.out_path_var.set(raw)
            return
        folder = os.path.dirname(raw)
        name = re.sub(r"\s*\(\d+\)$", "", os.path.splitext(os.path.basename(raw))[0])
        if name not in FORMATS:
            name = "35"
        ext = EXT_MAP.get(self.out_format_var.get(), ".xlsx")
        self.out_path_var.set(unique_path(folder, name, ext=ext))

    def pick_word(self):
        p = filedialog.askopenfilename(
            title="Выберите исходный файл",
            filetypes=[
                ("Все поддерживаемые", "*.docx *.doc *.odt *.ods"),
                ("Word 2007+ (.docx)", "*.docx"),
                ("Word 97-2003 (.doc)", "*.doc"),
                ("LibreOffice/OpenOffice текст (.odt)", "*.odt"),
                ("LibreOffice/OpenOffice таблица (.ods)", "*.ods"),
                ("Все файлы", "*.*"),
            ]
        )
        if p:
            self.word_path.set(p)

    def pick_save_path(self):
        folder = filedialog.askdirectory(title="Выберите папку для сохранения")
        if not folder:
            return
        self._save_folder = folder
        self.out_path_var.set(folder)

    def _set_buttons_enabled(self, enabled):
        state = "normal" if enabled else "disabled"
        for b in self.form_buttons:
            b.config(state=state)

    def run(self, fmt):
        w = self.word_path.get().strip()
        if not w:
            messagebox.showwarning("Не хватает данных", "Укажите исходный файл.")
            return

        folder = self._current_folder() or os.path.dirname(os.path.abspath(w))
        self._save_folder = folder

        out_path = self._show_full_target(fmt)

        out_fmt = self.out_format_var.get()
        sfo = self.sfo_var.get().strip()

        self._set_buttons_enabled(False)
        self._pause_idle()
        threading.Thread(
            target=self._worker,
            args=(w, fmt, out_fmt, out_path, sfo),
            daemon=True
        ).start()

    def _worker(self, w, fmt, out_fmt, out_path, sfo):
        try:
            result = do_transfer(
                w, fmt,
                out_format=out_fmt,
                out_path_manual=out_path,
                sfo_code=sfo,
            )
            self._set_buttons_enabled(True)
            self._show_folder_only()
            messagebox.showinfo("Готово", f"Файл сохранён:\n{result}")
        except Exception as ex:
            self._set_buttons_enabled(True)
            self._show_folder_only()
            messagebox.showerror("Ошибка", f"{ex}\n\n{traceback.format_exc()}")
        finally:
            self._resume_idle()


if __name__ == "__main__":
    root = tk.Tk()
    App(root)
    root.update_idletasks()
    root.deiconify()
    root.lift()
    root.attributes("-topmost", True)
    root.after(300, lambda: root.attributes("-topmost", False))
    root.focus_force()
    root.mainloop()
